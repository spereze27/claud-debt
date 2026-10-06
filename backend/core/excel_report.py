"""Exportación del linaje y los hallazgos a Excel (pandas + openpyxl).

Hojas:
  01_Inventario_Flujos   un flujo por fila: id, nombre, capas inferidas, fuente/destino crudos, grupo, orden…
  02_Dependencias_Flujos origen → destino (flujo que escribe un dato → flujo que lo lee)
  03_Lineaje             flujo ↔ dato (entradas y salidas) con sistema, capa, formato y modo de escritura
  04_Grupos_Cascada      resumen de cada grupo / cascada
  05_Hallazgos           hallazgos de todos los agentes
  00_Resumen             metadatos de la auditoría

Todas las hojas se crean aunque estén vacías (con sus columnas) para que el archivo tenga siempre la misma forma.
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from typing import Any

import pandas as pd

from backend.core.lineage import NOT_DETECTED

ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")
MAX_CELL = 32_000

INVENTORY_COLS = {
    "flujo_id": "ID", "nombre": "Nombre", "motor": "Motor", "proyecto": "Proyecto", "estado": "Estado",
    "capa": "Capa (entrada → salida)", "capa_entrada_inferida": "Capa entrada inferida",
    "capa_salida_inferida": "Capa salida inferida", "fuente_raw": "Fuente (raw)", "destino_raw": "Destino (raw)",
    "modo_escritura": "Modo de escritura", "grupo_cascada_id": "Grupo cascada", "origen_grupo": "Cómo se infirió el grupo",
    "orden_en_cascada": "Orden en cascada", "owner": "Owner", "calendario": "Calendario", "identidad": "Identidad",
    "ejecuciones": "Ejecuciones", "ultima_ejecucion": "Última ejecución", "duracion_media_s": "Duración media (s)",
    "duracion_p95_s": "Duración p95 (s)", "costo_30d_usd": "Costo 30 d (USD)", "observaciones": "Observaciones",
}
DEPENDENCY_COLS = ["origen_id", "origen_nombre", "destino_id", "destino_nombre", "dato", "capa_dato",
                   "modo_escritura_origen", "grupo_origen", "grupo_destino", "mismo_grupo"]
LINEAGE_COLS = ["flujo_id", "flujo", "grupo_cascada_id", "direccion", "dato_id", "dato_raw", "dato", "sistema", "capa",
                "formato", "modo_escritura"]
GROUP_COLS = ["grupo_cascada_id", "origen_inferencia", "num_flujos", "es_cascada", "profundidad", "flujos_en_orden",
              "capas_recorridas", "owners", "proyectos", "costo_30d_usd", "dependencias_con_otros_grupos", "estados"]
FINDING_COLS = ["severidad", "categoria", "regla", "recurso", "titulo", "descripcion", "recomendacion",
                "ahorro_mensual_usd"]


def clean_cell(v: Any) -> Any:
    """Texto seguro para Excel: sin caracteres de control, sin fórmulas, con longitud acotada. None → vacío."""
    if v is None:
        return ""
    if isinstance(v, bool) or isinstance(v, (int, float)):
        return v
    if isinstance(v, (list, tuple, set)):
        v = "; ".join(str(x) for x in v if x is not None)
    elif isinstance(v, dict):
        v = ", ".join(f"{k}={val}" for k, val in v.items())
    s = ILLEGAL_XML.sub("", str(v))
    if s.startswith(FORMULA_PREFIX):
        s = "'" + s
    return s[:MAX_CELL]


def _frame(rows: list[dict[str, Any]], columns: list[str]) -> pd.DataFrame:
    df = pd.DataFrame(rows, columns=columns)
    return df.map(clean_cell) if not df.empty else df


def _modes(outputs: list[dict[str, Any]]) -> str:
    modes = sorted({o.get("write_mode") for o in outputs if o.get("write_mode")})
    return ", ".join(modes) or NOT_DETECTED


def build_frames(state: dict[str, Any]) -> dict[str, pd.DataFrame]:
    raw = state.get("raw_infra_data") or {}
    dp = raw.get("data_platform") or {}
    flows = dp.get("flows") or []

    inventory = []
    for f in sorted(flows, key=lambda x: (x.get("grupo_cascada_id") or "~", x.get("orden_en_cascada") or 99, x.get("name") or "")):
        cin, cout = f.get("capa_entrada_inferida") or NOT_DETECTED, f.get("capa_salida_inferida") or NOT_DETECTED
        inventory.append({
            "flujo_id": f.get("resource_id"), "nombre": f.get("name"), "motor": f.get("engine"),
            "proyecto": f.get("project_id"), "estado": f.get("state_es") or f.get("state"),
            "capa": f"{cin} → {cout}", "capa_entrada_inferida": cin, "capa_salida_inferida": cout,
            "fuente_raw": f.get("fuente_raw") or NOT_DETECTED, "destino_raw": f.get("destino_raw") or NOT_DETECTED,
            "modo_escritura": _modes(f.get("outputs") or []), "grupo_cascada_id": f.get("grupo_cascada_id") or NOT_DETECTED,
            "origen_grupo": f.get("origen_grupo"), "orden_en_cascada": f.get("orden_en_cascada"),
            "owner": f.get("owner") or NOT_DETECTED, "calendario": f.get("schedule"), "identidad": f.get("identity"),
            "ejecuciones": f.get("runs_in_window"), "ultima_ejecucion": (f.get("last_run") or "")[:19],
            "duracion_media_s": f.get("avg_duration_s"), "duracion_p95_s": f.get("p95_duration_s"),
            "costo_30d_usd": f.get("cost_30d_usd"),
            "observaciones": " | ".join(s.get("text", "") for s in f.get("signals") or []),
        })
    inv = _frame(inventory, list(INVENTORY_COLS)).rename(columns=INVENTORY_COLS)

    lineage = []
    for f in flows:
        for direction, items in (("entrada", f.get("inputs") or []), ("salida", f.get("outputs") or [])):
            for i in items:
                lineage.append({"flujo_id": f.get("resource_id"), "flujo": f.get("name"),
                                "grupo_cascada_id": f.get("grupo_cascada_id"), "direccion": direction,
                                "dato_id": i.get("id"), "dato_raw": i.get("raw"), "dato": i.get("name"),
                                "sistema": i.get("system"), "capa": i.get("layer") or NOT_DETECTED,
                                "formato": i.get("format"), "modo_escritura": i.get("write_mode")})

    findings = [{"severidad": f.get("severity"), "categoria": f.get("category"), "regla": f.get("rule"),
                 "recurso": f.get("resource_id"), "titulo": f.get("title"), "descripcion": f.get("description"),
                 "recomendacion": f.get("recommendation"), "ahorro_mensual_usd": f.get("estimated_monthly_waste_usd")}
                for f in sorted(state.get("findings") or [],
                                key=lambda x: ["CRITICAL", "HIGH", "MEDIUM", "LOW"].index(x.get("severity", "LOW"))
                                if x.get("severity") in ("CRITICAL", "HIGH", "MEDIUM", "LOW") else 9)]

    k = state.get("kpis") or {}
    summary = [
        {"campo": "Proveedor", "valor": state.get("provider")}, {"campo": "Alcance", "valor": state.get("scope")},
        {"campo": "Objetivo", "valor": state.get("target_id")},
        {"campo": "Estado del descubrimiento", "valor": state.get("discovery_status")},
        {"campo": "Generado (UTC)", "valor": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")},
        {"campo": "Flujos", "valor": len(flows)}, {"campo": "Grupos / cascadas", "valor": len(dp.get("cascades") or [])},
        {"campo": "Dependencias entre flujos", "valor": len(dp.get("dependencies") or [])},
        {"campo": "Hallazgos", "valor": len(state.get("findings") or [])},
        {"campo": "Deuda técnica global (0-100)", "valor": k.get("overall_debt_score")},
        {"campo": "Calificación", "valor": k.get("grade")},
        {"campo": "Avisos de recolección", "valor": len(state.get("errors") or [])},
    ]
    return {
        "00_Resumen": _frame(summary, ["campo", "valor"]),
        "01_Inventario_Flujos": inv,
        "02_Dependencias_Flujos": _frame(dp.get("dependencies") or [], DEPENDENCY_COLS),
        "03_Lineaje": _frame(lineage, LINEAGE_COLS),
        "04_Grupos_Cascada": _frame(dp.get("cascades") or [], GROUP_COLS),
        "05_Hallazgos": _frame(findings, FINDING_COLS),
    }


def write_excel(frames: dict[str, pd.DataFrame], path: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in frames.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1A73E8")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for idx, col in enumerate(df.columns, start=1):
                sample = [len(str(col))] + [len(str(v)) for v in df[col].head(200)]
                ws.column_dimensions[get_column_letter(idx)].width = min(60, max(10, max(sample) + 2))
